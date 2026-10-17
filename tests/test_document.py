"""Tests for app.document module.
Verifies:
1. Multi-paragraph table cells are explicitly iterated via cell.paragraphs (never flattened or skipped).
2. Each paragraph inside table cells is extracted as an individual translatable segment.
3. Writing translated text preserves paragraph structures, runs, and bullet list styles (w:numPr).
4. Full reconstruction of real docx files with tables and multi-paragraph cells.
"""
import io
import unittest

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from app.document import (
    extract_segments_from_docx,
    is_translatable,
    iter_document_paragraphs,
    iter_table_paragraphs,
    reconstruct_docx,
    update_paragraph_text,
)


class TestDocumentProcessing(unittest.TestCase):
    def _create_sample_docx_with_multi_paragraph_table(self) -> bytes:
        doc = docx.Document()
        doc.add_heading("Document Title", level=1)
        doc.add_paragraph("This is a lead body paragraph.")

        # Create a 2x2 table
        table = doc.add_table(rows=2, cols=2)

        # Cell (0, 0): Normal single-paragraph cell
        cell_00 = table.cell(0, 0)
        cell_00.paragraphs[0].text = "Header 1"

        # Cell (0, 1): Normal single-paragraph cell
        cell_01 = table.cell(0, 1)
        cell_01.paragraphs[0].text = "Header 2"

        # Cell (1, 0): Simple cell
        cell_10 = table.cell(1, 0)
        cell_10.paragraphs[0].text = "Feature item"

        # Cell (1, 1): Multi-paragraph cell with bullet points and bold/italic styles
        cell_11 = table.cell(1, 1)
        p1 = cell_11.paragraphs[0]
        p1.style = "List Bullet"
        r1 = p1.add_run("Bullet point A: High performance")
        r1.bold = True

        p2 = cell_11.add_paragraph(style="List Bullet")
        r2 = p2.add_run("Bullet point B: Robust error isolation")
        r2.italic = True

        p3 = cell_11.add_paragraph(style="List Bullet")
        r3 = p3.add_run("Bullet point C: Automated retry")

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()

    def test_multi_paragraph_cells_not_skipped_or_flattened(self):
        """Verifies explicitly iterating cell.paragraphs yields all paragraphs without flattening."""
        docx_bytes = self._create_sample_docx_with_multi_paragraph_table()
        doc = docx.Document(io.BytesIO(docx_bytes))
        table = doc.tables[0]

        table_paras = list(iter_table_paragraphs(table))
        # Total paragraphs across table:
        # cell (0,0): 1
        # cell (0,1): 1
        # cell (1,0): 1
        # cell (1,1): 3
        # Total = 6 paragraphs
        self.assertEqual(len(table_paras), 6)

        texts = [p.text for p in table_paras]
        self.assertIn("Bullet point A: High performance", texts)
        self.assertIn("Bullet point B: Robust error isolation", texts)
        self.assertIn("Bullet point C: Automated retry", texts)

    def test_extract_segments_from_docx_individual_segments(self):
        """Each paragraph inside table cells is extracted as an individual translatable segment."""
        docx_bytes = self._create_sample_docx_with_multi_paragraph_table()
        segments = extract_segments_from_docx(docx_bytes)

        # Expected segments in document order:
        # 1. "Document Title"
        # 2. "This is a lead body paragraph."
        # 3. "Header 1"
        # 4. "Header 2"
        # 5. "Feature item"
        # 6. "Bullet point A: High performance"
        # 7. "Bullet point B: Robust error isolation"
        # 8. "Bullet point C: Automated retry"
        self.assertEqual(len(segments), 8)
        self.assertEqual(segments[0], "Document Title")
        self.assertEqual(segments[1], "This is a lead body paragraph.")
        self.assertEqual(segments[5], "Bullet point A: High performance")
        self.assertEqual(segments[6], "Bullet point B: Robust error isolation")
        self.assertEqual(segments[7], "Bullet point C: Automated retry")

    def test_update_paragraph_text_preserves_formatting_and_runs(self):
        """Updating paragraph text retains run styling (bold, italic) and paragraph style."""
        doc = docx.Document()
        p = doc.add_paragraph(style="List Bullet")
        r = p.add_run("Original text")
        r.bold = True

        update_paragraph_text(p, "Translated text")

        self.assertEqual(p.text, "Translated text")
        self.assertEqual(p.style.name, "List Bullet")
        self.assertTrue(p.runs[0].bold)

    def test_reconstruct_docx_multi_paragraph_cells(self):
        """Translating and reconstructing a docx updates each cell paragraph independently."""
        docx_bytes = self._create_sample_docx_with_multi_paragraph_table()
        translations = {
            "Document Title": "নথির শিরোনাম",
            "This is a lead body paragraph.": "এটি একটি প্রাথমিক মূল অনুচ্ছেদ।",
            "Header 1": "শিরোনাম ১",
            "Header 2": "শিরোনাম ২",
            "Feature item": "বৈশিষ্ট্য উপাদান",
            "Bullet point A: High performance": "বুলেট পয়েন্ট ক: উচ্চ কর্মক্ষমতা",
            "Bullet point B: Robust error isolation": "বুলেট পয়েন্ট খ: শক্তিশালী ত্রুটি বিচ্ছিন্নতা",
            "Bullet point C: Automated retry": "বুলেট পয়েন্ট গ: স্বয়ংক্রিয় পুনঃচেষ্টা",
        }

        reconstructed_bytes = reconstruct_docx(docx_bytes, translations)
        reconstructed_doc = docx.Document(io.BytesIO(reconstructed_bytes))

        # Check multi-paragraph cell (1, 1)
        multi_cell = reconstructed_doc.tables[0].cell(1, 1)
        self.assertEqual(len(multi_cell.paragraphs), 3)

        self.assertEqual(multi_cell.paragraphs[0].text, "বুলেট পয়েন্ট ক: উচ্চ কর্মক্ষমতা")
        self.assertEqual(multi_cell.paragraphs[0].style.name, "List Bullet")
        self.assertTrue(multi_cell.paragraphs[0].runs[0].bold)

        self.assertEqual(multi_cell.paragraphs[1].text, "বুলেট পয়েন্ট খ: শক্তিশালী ত্রুটি বিচ্ছিন্নতা")
        self.assertEqual(multi_cell.paragraphs[1].style.name, "List Bullet")
        self.assertTrue(multi_cell.paragraphs[1].runs[0].italic)

        self.assertEqual(multi_cell.paragraphs[2].text, "বুলেট পয়েন্ট গ: স্বয়ংক্রিয় পুনঃচেষ্টা")
        self.assertEqual(multi_cell.paragraphs[2].style.name, "List Bullet")

    def test_merged_cells_deduplication(self):
        """Merged cells are not processed repeatedly."""
        doc = docx.Document()
        table = doc.add_table(rows=2, cols=2)
        # Merge top row cells (0, 0) and (0, 1)
        cell_00 = table.cell(0, 0)
        cell_01 = table.cell(0, 1)
        cell_00.merge(cell_01)
        cell_00.paragraphs[0].text = "Merged Header"

        table.cell(1, 0).paragraphs[0].text = "Row 2 Col 1"
        table.cell(1, 1).paragraphs[0].text = "Row 2 Col 2"

        paras = list(iter_table_paragraphs(table))
        # Merged header should only appear once!
        merged_count = sum(1 for p in paras if p.text == "Merged Header")
        self.assertEqual(merged_count, 1)
        self.assertEqual(len(paras), 3)

    def test_nested_tables_traversal(self):
        """Nested tables within cells are recursively processed."""
        doc = docx.Document()
        outer_table = doc.add_table(rows=1, cols=1)
        cell = outer_table.cell(0, 0)
        cell.paragraphs[0].text = "Outer Cell Para"

        nested_table = cell.add_table(rows=1, cols=1)
        nested_cell = nested_table.cell(0, 0)
        nested_cell.paragraphs[0].text = "Nested Cell Para 1"
        nested_cell.add_paragraph("Nested Cell Para 2")

        paras = list(iter_table_paragraphs(outer_table))
        texts = [p.text for p in paras if p.text.strip()]
        self.assertEqual(texts, ["Outer Cell Para", "Nested Cell Para 1", "Nested Cell Para 2"])


if __name__ == "__main__":
    unittest.main()
